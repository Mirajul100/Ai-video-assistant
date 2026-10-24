import os
import zipfile

import docx
import pypdf
from pptx import Presentation
from langchain_core.documents import Document


SUPPORTED_DOCUMENT_EXTENSIONS = {".pdf", ".docx", ".pptx", ".txt", ".md"}
MAX_ARCHIVE_ENTRIES = 10_000
MAX_ARCHIVE_UNCOMPRESSED_BYTES = 250 * 1024 * 1024


class InvalidDocumentError(ValueError):
    """Raised when an uploaded document is invalid or unsupported."""


def _validate_text_file(file_path: str) -> None:
    try:
        with open(file_path, "r", encoding="utf-8-sig") as file:
            for chunk in iter(lambda: file.read(64 * 1024), ""):
                if "\x00" in chunk:
                    raise InvalidDocumentError("Text documents cannot contain binary data.")
    except UnicodeDecodeError as error:
        raise InvalidDocumentError("Text documents must use UTF-8 encoding.") from error


def _validate_ooxml_file(file_path: str, required_member: str) -> None:
    try:
        with zipfile.ZipFile(file_path) as archive:
            members = archive.infolist()
            if len(members) > MAX_ARCHIVE_ENTRIES:
                raise InvalidDocumentError("The document contains too many archive entries.")
            if sum(member.file_size for member in members) > MAX_ARCHIVE_UNCOMPRESSED_BYTES:
                raise InvalidDocumentError("The uncompressed document is too large.")
            if required_member not in archive.namelist():
                raise InvalidDocumentError("The file is not a valid Office document.")
            if archive.testzip() is not None:
                raise InvalidDocumentError("The Office document archive is corrupt.")
    except zipfile.BadZipFile as error:
        raise InvalidDocumentError("The file is not a valid Office document.") from error


def load_pdf_pages(file_path: str) -> list[Document]:
    try:
        reader = pypdf.PdfReader(file_path)
        if reader.is_encrypted:
            raise InvalidDocumentError("Encrypted PDFs are not supported.")
        return [
            Document(
                page_content=page.extract_text() or "",
                metadata={"source": file_path, "page": i},
            )
            for i, page in enumerate(reader.pages)
        ]
    except InvalidDocumentError:
        raise
    except Exception as error:
        raise InvalidDocumentError("The PDF is corrupt or could not be read.") from error


def load_word_pages(file_path: str) -> list[Document]:
    try:
        doc = docx.Document(file_path)
        full_text = [para.text for para in doc.paragraphs if para.text.strip()]
        return [Document(
            page_content="\n\n".join(full_text),
            metadata={"source": file_path}
        )]
    except Exception as error:
        raise InvalidDocumentError("The Word document is corrupt or could not be read.") from error


def load_ppt_pages(file_path: str) -> list[Document]:
    try:
        prs = Presentation(file_path)
        docs = []
        for i, slide in enumerate(prs.slides):
            slide_text = []
            for shape in slide.shapes:
                if hasattr(shape, "text") and shape.text.strip():
                    slide_text.append(shape.text.strip())
            if slide_text:
                docs.append(Document(
                    page_content="\n".join(slide_text),
                    metadata={"source": file_path, "slide": i + 1}
                ))
        return docs
    except Exception as error:
        raise InvalidDocumentError("The PowerPoint document is corrupt or could not be read.") from error


def load_text_pages(file_path: str) -> list[Document]:
    _validate_text_file(file_path)
    try:
        with open(file_path, "r", encoding="utf-8-sig") as file:
            text = file.read()
    except (OSError, UnicodeDecodeError) as error:
        raise InvalidDocumentError("The text document could not be read as UTF-8.") from error
    return [Document(
        page_content=text,
        metadata={"source": file_path}
    )]


def load_any_document(file_path: str) -> list[Document]:
    """Validate and load PDF, DOCX, PPTX, TXT, or Markdown documents."""
    if not os.path.isfile(file_path):
        raise InvalidDocumentError("The uploaded document could not be found.")

    ext = os.path.splitext(file_path)[1].lower()
    if ext not in SUPPORTED_DOCUMENT_EXTENSIONS:
        raise InvalidDocumentError("Unsupported document format.")

    try:
        if ext == ".pdf":
            with open(file_path, "rb") as file:
                if b"%PDF-" not in file.read(1024):
                    raise InvalidDocumentError("The file is not a valid PDF.")
            return load_pdf_pages(file_path)
        if ext == ".docx":
            _validate_ooxml_file(file_path, "word/document.xml")
            return load_word_pages(file_path)
        if ext == ".pptx":
            _validate_ooxml_file(file_path, "ppt/presentation.xml")
            return load_ppt_pages(file_path)
        return load_text_pages(file_path)
    except InvalidDocumentError:
        raise
    except OSError as error:
        raise InvalidDocumentError("The document could not be read.") from error
